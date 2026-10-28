
import json
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".json"}
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MIN_TEXT_LENGTH = 30
MAX_TEXT_LENGTH = 50_000


class JobDescriptionError(ValueError):
    """Raised when a job description cannot be extracted or validated."""


def _validate_text(text: str) -> str:
    """Normalize whitespace and validate extracted job description text."""
    text = " ".join(text.split())

    if len(text) < MIN_TEXT_LENGTH:
        raise JobDescriptionError(
            "Job description is too short or contains no readable text."
        )

    if len(text) > MAX_TEXT_LENGTH:
        raise JobDescriptionError(
            f"Job description exceeds {MAX_TEXT_LENGTH} characters."
        )

    return text


def _extract_json_text(raw_text: str) -> str:
    """Extract useful text from common JSON job-description structures."""
    try:
        data = json.loads(raw_text)
    except json.JSONDecodeError as exc:
        raise JobDescriptionError(
            "Invalid JSON. Please upload a valid JSON job description."
        ) from exc

    if isinstance(data, str):
        return data

    if isinstance(data, dict):
        # Prefer common job-description fields when present.
        preferred_fields = (
            "job_description",
            "description",
            "responsibilities",
            "requirements",
            "qualifications",
            "skills",
            "title",
        )

        parts = []
        for field in preferred_fields:
            value = data.get(field)
            if isinstance(value, str):
                parts.append(value)
            elif isinstance(value, list):
                parts.extend(str(item) for item in value)

        if parts:
            return "\n".join(parts)

    # For other valid JSON shapes, retain the data instead of losing it.
    return json.dumps(data, ensure_ascii=False)


def parse_job_description(
    filename: str,
    file_bytes: bytes,
) -> str:
    """
    Extract and validate job-description text from PDF, DOCX, TXT or JSON.

    Returns normalized text ready for the matching pipeline.
    """
    if not file_bytes:
        raise JobDescriptionError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise JobDescriptionError("File exceeds the 5 MB upload limit.")

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise JobDescriptionError(
            "Unsupported file type. Upload PDF, DOCX, TXT or JSON."
        )

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(file_bytes))

            if reader.is_encrypted:
                raise JobDescriptionError(
                    "Encrypted PDFs are not supported. Upload an unlocked PDF."
                )

            extracted = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )

        elif extension == ".docx":
            document = Document(BytesIO(file_bytes))
            paragraphs = [p.text for p in document.paragraphs]

            # Include text in tables, where job requirements are often listed.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            extracted = "\n".join(paragraphs)

        elif extension == ".txt":
            extracted = file_bytes.decode("utf-8-sig")

        else:  # JSON
            raw_text = file_bytes.decode("utf-8-sig")
            extracted = _extract_json_text(raw_text)

    except JobDescriptionError:
        raise
    except Exception as exc:
        raise JobDescriptionError(
            "Could not read this file. Check that its contents match its extension."
        ) from exc

    return _validate_text(extracted)


def parse_job_description_text(text: str) -> str:
    """Validate a job description pasted directly into the application."""
    if not isinstance(text, str):
        raise JobDescriptionError("Job description must be text.")

    return _validate_text(text)
